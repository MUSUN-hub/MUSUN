import tkinter as tk
from tkinter import filedialog, messagebox
import os
import json
import threading
import subprocess
import urllib.request
import urllib.parse
import sys
import platform

# ── 라이선스 서버 ─────────────────────────────────────────
LICENSE_SERVER_URL = "https://script.google.com/macros/s/AKfycbzx9K_0KmVZph-O_MyEXAH7AbKWOwL5zu5Z3_re2b2w6PoZJnClGPipoU9Wah_dhNaL/exec"
if platform.system() == "Darwin":
    LICENSE_PATH = os.path.join(os.path.expanduser("~"), "Library", "Application Support", "PPTAnalyzer", "license.key")
else:
    LICENSE_PATH = os.path.join(os.environ.get("APPDATA", ""), "PPTAnalyzer", "license.key")

def generate_hwid():
    try:
        if platform.system() == "Darwin":
            result = subprocess.check_output(
                ["ioreg", "-rd1", "-c", "IOPlatformExpertDevice"],
                timeout=5
            ).decode()
            for line in result.splitlines():
                if "IOPlatformUUID" in line:
                    uuid = line.split('"')[-2]
                    if uuid and len(uuid) > 10:
                        return uuid.upper()
        else:
            result = subprocess.check_output(
                'powershell -NoProfile -Command "(Get-WmiObject -Class Win32_ComputerSystemProduct).UUID"',
                shell=True, timeout=5
            ).decode().strip()
            if result and len(result) > 10:
                return result.upper()
    except Exception:
        pass
    # fallback
    import socket
    hostname = socket.gethostname()
    username = os.environ.get("USERNAME", os.environ.get("USER", "UNKNOWN"))
    return f"FALLBACK_{hostname}_{username}"

def server_request(action, params=None):
    try:
        qs = {"action": action}
        if params:
            qs.update(params)
        url = LICENSE_SERVER_URL + "?" + urllib.parse.urlencode(qs)
        with urllib.request.urlopen(url, timeout=10) as resp:
            return json.loads(resp.read().decode())
    except Exception:
        return None

def load_local_license():
    try:
        if not os.path.exists(LICENSE_PATH):
            return None
        with open(LICENSE_PATH, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return None

def save_local_license(hwid, key=None, email=None, last_verified=None, expiry_date="", plan="",
                       notified_30d=False, notified_7d=False):
    try:
        import time
        os.makedirs(os.path.dirname(LICENSE_PATH), exist_ok=True)
        with open(LICENSE_PATH, "w", encoding="utf-8") as f:
            json.dump({
                "hwid": hwid,
                "key": key,
                "email": email,
                "last_verified": last_verified or time.time(),
                "expiry_date": expiry_date,
                "plan": plan,
                "notified_30d": notified_30d,
                "notified_7d": notified_7d
            }, f)
        return True
    except Exception:
        return False

def get_expiry_display(expiry_date, plan):
    if not expiry_date:
        return ""
    try:
        from datetime import datetime, timezone
        expiry = datetime.fromisoformat(expiry_date.replace("Z", "+00:00"))
        if expiry.tzinfo is None:
            expiry = expiry.replace(tzinfo=timezone.utc)
        now = datetime.now(timezone.utc)
        days = (expiry - now).days
        plan_str = "월간" if plan == "monthly" else "연간"
        date_str = expiry.strftime("%Y-%m-%d")
        if days < 0:
            grace_left = 7 + days
            if grace_left > 0:
                return f"⚠  구독 만료 — 유예기간 {grace_left}일 남음 ({date_str})"
            else:
                return f"✕  구독 만료 ({date_str})"
        elif days <= 7:
            return f"⚠  {plan_str} 구독 — 만료 {days}일 전 ({date_str})"
        elif days <= 30:
            return f"△  {plan_str} 구독 — 만료 {days}일 전 ({date_str})"
        else:
            return f"{plan_str} 구독  |  만료일: {date_str}"
    except Exception:
        return ""

def delete_local_license():
    try:
        if os.path.exists(LICENSE_PATH):
            os.remove(LICENSE_PATH)
    except Exception:
        pass

# ── 폰트 ────────────────────────────────────────────────
_SYS_FONT = "Apple SD Gothic Neo" if platform.system() == "Darwin" else "Malgun Gothic"

# ── 컬러 팔레트 ──────────────────────────────────────────
BG_DEEP    = "#0e0e12"   # 최상위 배경
BG_CARD    = "#16161d"   # 카드 배경
BG_INPUT   = "#1e1e28"   # 입력창 배경
BG_BORDER  = "#2a2a38"   # 테두리
ACCENT     = "#c9a85c"   # 골드 액센트
ACCENT_DIM = "#8a7040"   # 골드 딤
FG_PRIMARY = "#e8e6e0"   # 주 텍스트
FG_MUTED   = "#7a7870"   # 보조 텍스트
FG_LOG     = "#9db8a0"   # 로그 텍스트
BTN_BG     = "#c9a85c"   # 버튼
BTN_FG     = "#0e0e12"   # 버튼 텍스트
BTN_HOV    = "#e0bf70"   # 버튼 호버
# ─────────────────────────────────────────────────────────

class LicenseDialog:
    """라이선스 체크 다이얼로그 — 앱 시작 시 호출"""

    def __init__(self, root):
        self.root = root
        self.hwid = generate_hwid()
        self.result = False
        self.expiry_date = ""
        self.plan = ""
        self.license_info = ""

    def check_and_run(self):
        import time
        VERIFY_INTERVAL = 7 * 24 * 3600

        local = load_local_license()

        # ① 라이선스 키 보유 → 구독 확인
        if local and local.get("hwid") == self.hwid and local.get("key"):
            elapsed = time.time() - local.get("last_verified", 0)
            if elapsed < VERIFY_INTERVAL:
                self.expiry_date = local.get("expiry_date", "")
                self.plan = local.get("plan", "")
                self._check_expiry_and_proceed()
                return
            resp = server_request("check_subscription", {"hwid": self.hwid, "key": local.get("key")})
            if resp is None:
                self.expiry_date = local.get("expiry_date", "")
                self.plan = local.get("plan", "")
                self._check_expiry_and_proceed()
                return
            status = resp.get("status", "")
            if status == "active":
                self.expiry_date = resp.get("expiry_date", "")
                self.plan = resp.get("plan", "")
                save_local_license(self.hwid, key=local.get("key"),
                                   expiry_date=self.expiry_date, plan=self.plan)
                self._check_expiry_and_proceed()
                return
            elif status == "grace_period":
                self.expiry_date = resp.get("expiry_date", "")
                self.plan = resp.get("plan", "")
                save_local_license(self.hwid, key=local.get("key"),
                                   expiry_date=self.expiry_date, plan=self.plan)
                messagebox.showwarning("구독 만료",
                    f"구독이 만료되었습니다.\n유예기간: {resp.get('grace_days_left', 0)}일 남았습니다.\n갱신해주세요.")
                self.result = True
                return
            elif status == "expired":
                delete_local_license()
                messagebox.showerror("구독 만료", "구독이 만료되었습니다.\n라이선스 키를 다시 입력해주세요.")
                self._show_key_input()
                return
            else:
                self.expiry_date = local.get("expiry_date", "")
                self.plan = local.get("plan", "")
                self._check_expiry_and_proceed()
                return

        # ② 이메일(체험판) 보유 → trial_check
        if local and local.get("email"):
            resp = server_request("trial_check", {"hwid": self.hwid, "email": local.get("email")})
            if resp is None:
                self.license_info = "체험판 (오프라인)"
                self.result = True
                return
            status = resp.get("status", "")
            if status in ("trial_active", "trial_already_active"):
                remaining = resp.get("remaining", 0)
                self.license_info = f"체험판  |  남은 기간: {remaining}일"
                self._show_trial_active(remaining)
                return
            elif status == "trial_expired":
                self._show_expired_dialog()
                return
            else:
                # 알 수 없는 상태 → 웰컴 다이얼로그로
                pass

        # ③ 아무것도 없음 → 웰컴 (체험 시작 or 키 입력)
        self._show_start_dialog()

    def _check_expiry_and_proceed(self):
        if not self.expiry_date:
            self.license_info = "라이선스 활성"
            self.result = True
            return
        try:
            from datetime import datetime, timezone
            expiry = datetime.fromisoformat(self.expiry_date.replace("Z", "+00:00"))
            if expiry.tzinfo is None:
                expiry = expiry.replace(tzinfo=timezone.utc)
            now = datetime.now(timezone.utc)
            days = (expiry - now).days
            is_yearly = self.plan == "yearly"

            local = load_local_license()
            notified_30d = local.get("notified_30d", False) if local else False
            notified_7d  = local.get("notified_7d",  False) if local else False

            if days > 30:
                self.result = True
            elif days > 7:
                if is_yearly and not notified_30d:
                    messagebox.showinfo("구독 갱신 안내",
                        f"구독 만료까지 {days}일 남았습니다.\n기간 내 갱신해주세요.")
                    save_local_license(local["hwid"], key=local.get("key"), email=local.get("email"),
                                       last_verified=local.get("last_verified"),
                                       expiry_date=self.expiry_date, plan=self.plan,
                                       notified_30d=True, notified_7d=notified_7d)
                self.result = True
            elif days > 0:
                if not notified_7d:
                    messagebox.showwarning("구독 만료 임박",
                        f"구독 만료까지 {days}일 남았습니다!\n빠른 시일 내 갱신해주세요.")
                    save_local_license(local["hwid"], key=local.get("key"), email=local.get("email"),
                                       last_verified=local.get("last_verified"),
                                       expiry_date=self.expiry_date, plan=self.plan,
                                       notified_30d=notified_30d, notified_7d=True)
                self.result = True
            elif days >= -7:
                grace_left = 7 + days
                messagebox.showwarning("구독 만료 — 유예기간",
                    f"구독이 만료되었습니다.\n유예기간: {grace_left}일 남았습니다.\n갱신해주세요.")
                self.result = True
            else:
                messagebox.showerror("구독 만료",
                    "구독 유예기간이 종료되었습니다.\n라이선스 키를 다시 입력해주세요.")
                delete_local_license()
                self._show_key_input()
        except Exception:
            self.result = True

    def _show_start_dialog(self):
        """처음 실행 — 체험판 시작(이메일 인증) or 라이선스 키 입력"""
        win = tk.Toplevel(self.root)
        win.title("PPT Slide Analyzer")
        win.geometry("400x300")
        win.configure(bg=BG_DEEP)
        win.resizable(False, False)
        win.grab_set()

        tk.Label(win, text="PPT SLIDE ANALYZER",
                 font=(_SYS_FONT, 14, "bold"), fg=ACCENT, bg=BG_DEEP).pack(pady=(28, 4))
        tk.Label(win, text="이 소프트웨어는 라이선스가 필요합니다.",
                 font=(_SYS_FONT, 9), fg=FG_MUTED, bg=BG_DEEP).pack()
        tk.Label(win, text="7일 무료 체험판을 시작하거나 라이선스 키를 입력하세요.",
                 font=(_SYS_FONT, 9), fg=FG_MUTED, bg=BG_DEEP).pack(pady=(2, 16))

        btn_frame = tk.Frame(win, bg=BG_DEEP)
        btn_frame.pack(fill="x", padx=40)

        tk.Button(btn_frame, text="7일 무료 체험 시작",
                  command=lambda: [win.destroy(), self._show_email_step()],
                  bg=BTN_BG, fg=BTN_FG, relief="flat",
                  font=(_SYS_FONT, 10, "bold"), cursor="hand2", pady=9
                  ).pack(fill="x", pady=(0, 8))

        tk.Button(btn_frame, text="라이선스 키 입력",
                  command=lambda: [win.destroy(), self._show_key_input()],
                  bg=BG_CARD, fg=FG_PRIMARY, relief="flat",
                  font=(_SYS_FONT, 10), cursor="hand2", pady=9,
                  highlightbackground=BG_BORDER, highlightthickness=1
                  ).pack(fill="x")

        win.wait_window()

    def _show_email_step(self):
        """Step 2 — 이메일 입력 → 인증 코드 발송"""
        win = tk.Toplevel(self.root)
        win.title("체험판 시작 — 이메일 인증")
        win.geometry("420x230")
        win.configure(bg=BG_DEEP)
        win.resizable(False, False)
        win.grab_set()

        tk.Label(win, text="이메일 주소 입력",
                 font=(_SYS_FONT, 12, "bold"), fg=FG_PRIMARY, bg=BG_DEEP).pack(pady=(28, 6))
        tk.Label(win, text="입력하신 이메일로 인증 코드를 발송합니다.\n이메일당 1회만 체험판 사용 가능합니다.",
                 font=(_SYS_FONT, 9), fg=FG_MUTED, bg=BG_DEEP).pack(pady=(0, 12))

        email_var = tk.StringVar()
        entry = tk.Entry(win, textvariable=email_var, font=(_SYS_FONT, 11),
                         bg=BG_INPUT, fg=FG_PRIMARY, insertbackground=ACCENT,
                         relief="flat", highlightbackground=BG_BORDER, highlightthickness=1,
                         justify="center")
        entry.pack(fill="x", padx=40, ipady=7)
        entry.focus()

        msg_label = tk.Label(win, text="", font=(_SYS_FONT, 9), bg=BG_DEEP, fg=FG_MUTED)
        msg_label.pack(pady=(6, 0))

        def send_code():
            email = email_var.get().strip()
            if not email or "@" not in email:
                msg_label.config(text="올바른 이메일 주소를 입력하세요.", fg="#e06060")
                return
            msg_label.config(text="코드 발송 중...", fg=FG_MUTED)
            win.update()
            resp = server_request("trial_request_code", {"hwid": self.hwid, "email": email})
            if resp is None:
                msg_label.config(text="서버 연결 실패. 인터넷을 확인하세요.", fg="#e06060")
                return
            status = resp.get("status", "")
            if status == "code_sent":
                win.destroy()
                self._show_code_step(email)
            elif status == "trial_already_active":
                save_local_license(self.hwid, email=email)
                win.destroy()
                remaining = resp.get("remaining", 7)
                self.license_info = f"체험판  |  남은 기간: {remaining}일"
                messagebox.showinfo("체험판 활성", f"체험판이 활성화되었습니다.\n남은 기간: {remaining}일")
                self.result = True
            elif status == "trial_email_used":
                msg_label.config(text="이미 다른 기기에서 사용된 이메일입니다.", fg="#e06060")
            elif status == "code_already_sent":
                win.destroy()
                self._show_code_step(email)
            else:
                msg_label.config(text=f"오류: {resp.get('message', status)}", fg="#e06060")

        entry.bind("<Return>", lambda e: send_code())
        tk.Button(win, text="인증 코드 발송", command=send_code,
                  bg=BTN_BG, fg=BTN_FG, relief="flat",
                  font=(_SYS_FONT, 10, "bold"), cursor="hand2", pady=8
                  ).pack(fill="x", padx=40, pady=(12, 0))

        win.wait_window()

    def _show_code_step(self, email):
        """Step 3 — 인증 코드 입력"""
        win = tk.Toplevel(self.root)
        win.title("인증 코드 입력")
        win.geometry("420x240")
        win.configure(bg=BG_DEEP)
        win.resizable(False, False)
        win.grab_set()

        tk.Label(win, text="인증 코드 입력",
                 font=(_SYS_FONT, 12, "bold"), fg=FG_PRIMARY, bg=BG_DEEP).pack(pady=(28, 4))
        tk.Label(win, text=f"{email}\n으로 발송된 6자리 코드를 입력하세요.",
                 font=(_SYS_FONT, 9), fg=FG_MUTED, bg=BG_DEEP).pack(pady=(0, 12))

        code_var = tk.StringVar()
        entry = tk.Entry(win, textvariable=code_var, font=(_SYS_FONT, 14),
                         bg=BG_INPUT, fg=FG_PRIMARY, insertbackground=ACCENT,
                         relief="flat", highlightbackground=BG_BORDER, highlightthickness=1,
                         justify="center", width=12)
        entry.pack(ipady=7)
        entry.focus()

        msg_label = tk.Label(win, text="", font=(_SYS_FONT, 9), bg=BG_DEEP, fg=FG_MUTED)
        msg_label.pack(pady=(6, 0))

        def verify():
            code = code_var.get().strip()
            if not code:
                msg_label.config(text="코드를 입력하세요.", fg="#e06060")
                return
            msg_label.config(text="확인 중...", fg=FG_MUTED)
            win.update()
            resp = server_request("trial_verify_code", {"hwid": self.hwid, "email": email, "code": code})
            if resp is None:
                msg_label.config(text="서버 연결 실패. 인터넷을 확인하세요.", fg="#e06060")
                return
            status = resp.get("status", "")
            if status in ("verified", "trial_already_active"):
                save_local_license(self.hwid, email=email)
                remaining = resp.get("remaining", 7)
                self.license_info = f"체험판  |  남은 기간: {remaining}일"
                win.destroy()
                messagebox.showinfo("체험판 시작", f"체험판이 시작되었습니다.\n남은 기간: {remaining}일")
                self.result = True
            elif status == "invalid_code":
                msg_label.config(text="올바르지 않은 코드입니다.", fg="#e06060")
            elif status == "code_expired":
                msg_label.config(text="코드가 만료되었습니다. 다시 시도해주세요.", fg="#e06060")
            else:
                msg_label.config(text=f"오류: {resp.get('message', status)}", fg="#e06060")

        entry.bind("<Return>", lambda e: verify())
        tk.Button(win, text="확인", command=verify,
                  bg=BTN_BG, fg=BTN_FG, relief="flat",
                  font=(_SYS_FONT, 10, "bold"), cursor="hand2", pady=8
                  ).pack(fill="x", padx=40, pady=(12, 0))

        win.wait_window()

    def _show_trial_active(self, remaining):
        win = tk.Toplevel(self.root)
        win.title("PPT Slide Analyzer — 체험판")
        win.geometry("400x240")
        win.configure(bg=BG_DEEP)
        win.resizable(False, False)
        win.grab_set()

        tk.Label(win, text="PPT SLIDE ANALYZER",
                 font=(_SYS_FONT, 14, "bold"), fg=ACCENT, bg=BG_DEEP).pack(pady=(30, 4))
        tk.Label(win, text=f"체험판 사용 중  |  남은 기간: {remaining}일",
                 font=(_SYS_FONT, 10), fg=FG_PRIMARY, bg=BG_DEEP).pack(pady=(4, 20))

        btn_frame = tk.Frame(win, bg=BG_DEEP)
        btn_frame.pack(fill="x", padx=40)

        tk.Button(btn_frame, text="계속 체험판으로 사용",
                  command=lambda: [win.destroy(), setattr(self, 'result', True)],
                  bg=BTN_BG, fg=BTN_FG, relief="flat",
                  font=(_SYS_FONT, 10, "bold"), cursor="hand2", pady=9
                  ).pack(fill="x", pady=(0, 8))

        tk.Button(btn_frame, text="라이선스 키 등록",
                  command=lambda: [win.destroy(), self._show_key_input()],
                  bg=BG_CARD, fg=FG_PRIMARY, relief="flat",
                  font=(_SYS_FONT, 10), cursor="hand2", pady=9,
                  highlightbackground=BG_BORDER, highlightthickness=1
                  ).pack(fill="x")

        win.wait_window()

    def _show_expired_dialog(self):
        win = tk.Toplevel(self.root)
        win.title("체험 기간 만료")
        win.geometry("400x220")
        win.configure(bg=BG_DEEP)
        win.resizable(False, False)
        win.grab_set()

        tk.Label(win, text="체험 기간이 종료되었습니다.",
                 font=(_SYS_FONT, 13, "bold"), fg=ACCENT, bg=BG_DEEP).pack(pady=(36, 6))
        tk.Label(win, text="계속 사용하려면 라이선스 키를 등록하세요.",
                 font=(_SYS_FONT, 9), fg=FG_MUTED, bg=BG_DEEP).pack(pady=(0, 24))

        tk.Button(win, text="라이선스 키 입력",
                  command=lambda: [win.destroy(), self._show_key_input()],
                  bg=BTN_BG, fg=BTN_FG, relief="flat",
                  font=(_SYS_FONT, 10, "bold"), cursor="hand2", pady=9
                  ).pack(fill="x", padx=40)

        win.wait_window()

    def _show_key_input(self):
        win = tk.Toplevel(self.root)
        win.title("라이선스 키 입력")
        win.geometry("420x200")
        win.configure(bg=BG_DEEP)
        win.resizable(False, False)
        win.grab_set()

        tk.Label(win, text="라이선스 키를 입력하세요.",
                 font=(_SYS_FONT, 11, "bold"), fg=FG_PRIMARY, bg=BG_DEEP).pack(pady=(30, 12))

        key_var = tk.StringVar()
        entry = tk.Entry(win, textvariable=key_var, font=(_SYS_FONT, 11),
                         bg=BG_INPUT, fg=FG_PRIMARY, insertbackground=ACCENT,
                         relief="flat", highlightbackground=BG_BORDER, highlightthickness=1,
                         justify="center")
        entry.pack(fill="x", padx=40, ipady=7)
        entry.focus()

        msg_label = tk.Label(win, text="", font=(_SYS_FONT, 9), bg=BG_DEEP, fg=FG_MUTED)
        msg_label.pack(pady=(6, 0))

        def activate():
            key = key_var.get().strip()
            if not key:
                msg_label.config(text="키를 입력하세요.", fg=FG_MUTED)
                return
            msg_label.config(text="확인 중...", fg=FG_MUTED)
            win.update()
            resp = server_request("activate", {"hwid": self.hwid, "key": key})
            if resp is None:
                msg_label.config(text="서버 연결 실패. 인터넷을 확인하세요.", fg="#e06060")
                return
            status = resp.get("status", "")
            if status in ("activated", "already_active"):
                self.expiry_date = resp.get("expiry_date", "")
                self.plan = resp.get("plan", "")
                save_local_license(self.hwid, key=key,
                                   expiry_date=self.expiry_date, plan=self.plan)
                win.destroy()
                messagebox.showinfo("인증 완료", "라이선스가 등록되었습니다.")
                self.result = True
            elif status == "invalid_key":
                msg_label.config(text="유효하지 않은 키입니다.", fg="#e06060")
            elif status == "max_reached":
                msg_label.config(text=f"최대 등록 기기 수 초과 ({resp.get('max', 1)}대).", fg="#e06060")
            else:
                msg_label.config(text=f"오류: {status}", fg="#e06060")

        entry.bind("<Return>", lambda e: activate())
        tk.Button(win, text="확인", command=activate,
                  bg=BTN_BG, fg=BTN_FG, relief="flat",
                  font=(_SYS_FONT, 10, "bold"), cursor="hand2", pady=8
                  ).pack(fill="x", padx=40, pady=(12, 0))

        win.wait_window()


class PPTAnalyzerGUI:
    def __init__(self, root, expiry_date="", plan="", license_info=""):
        self.root = root
        self.expiry_date = expiry_date
        self.plan = plan
        self.license_info = license_info
        self.root.title("PPT Slide Analyzer")
        self.root.geometry("880x880")
        self.root.configure(bg=BG_DEEP)
        self.root.resizable(True, True)

        self.font_title  = (_SYS_FONT, 14, "bold")
        self.font_sub    = (_SYS_FONT,  9)
        self.font_label  = (_SYS_FONT, 10)
        self.font_log    = (_SYS_FONT,  9, "normal")
        self.font_btn    = (_SYS_FONT, 11, "bold")
        self.font_small  = (_SYS_FONT,  8)

        self._build_ui()

    # ── 유틸: 구분선 ────────────────────────────────────────
    def _divider(self, parent):
        tk.Frame(parent, bg=BG_BORDER, height=1).pack(fill="x", pady=(0, 16))

    # ── 카드 프레임 ─────────────────────────────────────────
    def _card(self, parent, title, pady_top=0):
        wrap = tk.Frame(parent, bg=BG_DEEP)
        wrap.pack(fill="x", pady=(pady_top, 18))

        # 라벨 행
        hdr = tk.Frame(wrap, bg=BG_DEEP)
        hdr.pack(fill="x", pady=(0, 8))
        tk.Label(hdr, text="▎", font=(_SYS_FONT, 12), fg=ACCENT,
                 bg=BG_DEEP).pack(side="left")
        tk.Label(hdr, text=title, font=self.font_label,
                 fg=FG_PRIMARY, bg=BG_DEEP).pack(side="left", padx=(4, 0))

        # 카드 본체
        body = tk.Frame(wrap, bg=BG_CARD, padx=18, pady=14,
                        highlightbackground=BG_BORDER, highlightthickness=1)
        body.pack(fill="x")
        return body

    # ── 스타일 입력창 ────────────────────────────────────────
    def _entry(self, parent, textvariable, width=None, **kw):
        e = tk.Entry(parent, textvariable=textvariable,
                     bg=BG_INPUT, fg=FG_PRIMARY, insertbackground=ACCENT,
                     relief="flat", font=self.font_label,
                     highlightbackground=BG_BORDER, highlightthickness=1,
                     **({} if width is None else {"width": width}), **kw)
        return e

    # ── 스타일 버튼 ─────────────────────────────────────────
    def _button(self, parent, text, command, width=None, pady=10, font=None):
        btn = tk.Button(
            parent, text=text, command=command,
            bg=BTN_BG, fg=BTN_FG,
            activebackground=BTN_HOV, activeforeground=BTN_FG,
            relief="flat", cursor="hand2",
            font=font or self.font_btn,
            pady=pady,
            **({"width": width} if width else {})
        )
        btn.bind("<Enter>", lambda e: btn.config(bg=BTN_HOV))
        btn.bind("<Leave>", lambda e: btn.config(bg=btn._default_bg if hasattr(btn, "_default_bg") else BTN_BG))
        btn._default_bg = BTN_BG
        return btn

    # ── UI 빌드 ─────────────────────────────────────────────
    def _build_ui(self):
        # ── 헤더 ────────────────────────────────────────────
        hdr = tk.Frame(self.root, bg=BG_CARD,
                       highlightbackground=BG_BORDER, highlightthickness=1)
        hdr.pack(fill="x")

        inner_hdr = tk.Frame(hdr, bg=BG_CARD, padx=32, pady=20)
        inner_hdr.pack(fill="x")

        tk.Label(inner_hdr, text="SLIDE MASTER ANALYZER",
                 font=(_SYS_FONT, 16, "bold"),
                 fg=ACCENT, bg=BG_CARD).pack(anchor="w")
        tk.Label(inner_hdr,
                 text="슬라이드 마스터 도형 텍스트 기반 강사 위치 감지",
                 font=self.font_sub, fg=FG_MUTED, bg=BG_CARD).pack(anchor="w", pady=(3, 0))

        expiry_text = get_expiry_display(self.expiry_date, self.plan) or self.license_info
        if expiry_text:
            fg = "#e06060" if "✕" in expiry_text or ("만료" in expiry_text and "유예" in expiry_text) \
                 else "#e0a030" if "⚠" in expiry_text or "△" in expiry_text \
                 else FG_MUTED
            tk.Label(inner_hdr, text=expiry_text,
                     font=self.font_small, fg=fg, bg=BG_CARD).pack(anchor="w", pady=(4, 0))

        # ── 본문 ────────────────────────────────────────────
        main = tk.Frame(self.root, bg=BG_DEEP, padx=30, pady=24)
        main.pack(fill="both", expand=True)

        # ① 파일 선택
        card1 = self._card(main, "PPT 파일 선택")
        file_row = tk.Frame(card1, bg=BG_CARD)
        file_row.pack(fill="x")

        self.ppt_path_var = tk.StringVar()
        path_entry = self._entry(file_row, self.ppt_path_var)
        path_entry.pack(side="left", fill="x", expand=True, ipady=6)

        browse_btn = self._button(file_row, "  열기  ", self._browse, pady=6,
                                  font=(_SYS_FONT, 10, "bold"))
        browse_btn.pack(side="left", padx=(10, 0))

        # ② 분석 설정
        card2 = self._card(main, "분석 설정")

        range_row = tk.Frame(card2, bg=BG_CARD)
        range_row.pack(fill="x", pady=(0, 12))

        tk.Label(range_row, text="슬라이드 범위",
                 font=self.font_label, fg=FG_MUTED, bg=BG_CARD, width=13, anchor="w").pack(side="left")

        self.start_slide_var = tk.StringVar(value="1")
        self._entry(range_row, self.start_slide_var, width=6).pack(side="left", ipady=5)

        tk.Label(range_row, text=" — ", font=self.font_label,
                 fg=FG_MUTED, bg=BG_CARD).pack(side="left")

        self.end_slide_var = tk.StringVar(value="999")
        self._entry(range_row, self.end_slide_var, width=6).pack(side="left", ipady=5)

        tk.Label(range_row, text="  (999 = 끝까지)",
                 font=self.font_small, fg=FG_MUTED, bg=BG_CARD).pack(side="left", padx=(6, 0))

        kw_row = tk.Frame(card2, bg=BG_CARD)
        kw_row.pack(fill="x")

        tk.Label(kw_row, text="감지 키워드",
                 font=self.font_label, fg=FG_MUTED, bg=BG_CARD, width=13, anchor="w").pack(side="left")

        self.shape_keyword_var = tk.StringVar(value="강사,교수자,instructor,teacher,유튜버")
        self._entry(kw_row, self.shape_keyword_var).pack(side="left", fill="x", expand=True, ipady=5)

        tk.Label(kw_row, text="  쉼표 구분",
                 font=self.font_small, fg=FG_MUTED, bg=BG_CARD).pack(side="left", padx=(6, 0))

        # ── 하단: 통계 + 버튼 ───────────────────────────────
        bottom = tk.Frame(main, bg=BG_DEEP)
        bottom.pack(fill="x", pady=(0, 12))

        self.stats_label = tk.Label(
            bottom,
            text="마스터/레이아웃 도형 텍스트로 강사 위치를 감지합니다.",
            font=self.font_small, fg=FG_MUTED, bg=BG_DEEP
        )
        self.stats_label.pack(anchor="w", pady=(0, 8))

        self.start_btn = self._button(bottom, "분석 시작", self._start_thread)
        self.start_btn.pack(fill="x", ipady=2)

        # ③ 진행 로그
        log_wrap = tk.Frame(main, bg=BG_DEEP)
        log_wrap.pack(fill="both", expand=True)

        log_hdr = tk.Frame(log_wrap, bg=BG_DEEP)
        log_hdr.pack(fill="x", pady=(0, 8))
        tk.Label(log_hdr, text="▎", font=(_SYS_FONT, 12),
                 fg=ACCENT, bg=BG_DEEP).pack(side="left")
        tk.Label(log_hdr, text="진행 로그", font=self.font_label,
                 fg=FG_PRIMARY, bg=BG_DEEP).pack(side="left", padx=(4, 0))

        log_card = tk.Frame(log_wrap, bg=BG_CARD,
                            highlightbackground=BG_BORDER, highlightthickness=1)
        log_card.pack(fill="both", expand=True)

        self.log_text = tk.Text(
            log_card, font=self.font_log, bg=BG_CARD,
            fg=FG_LOG, relief="flat", padx=16, pady=12,
            insertbackground=ACCENT, selectbackground=ACCENT_DIM,
            wrap="word"
        )
        scroll = tk.Scrollbar(log_card, orient="vertical",
                              command=self.log_text.yview, bg=BG_CARD,
                              troughcolor=BG_CARD, relief="flat")
        self.log_text.configure(yscrollcommand=scroll.set)
        scroll.pack(side="right", fill="y")
        self.log_text.pack(fill="both", expand=True)

    # ── 로그 출력 ───────────────────────────────────────────
    def _log(self, msg):
        self.log_text.insert(tk.END, msg + "\n")
        self.log_text.see(tk.END)
        self.root.update_idletasks()

    def _browse(self):
        f = filedialog.askopenfilename(filetypes=[("PowerPoint", "*.pptx *.ppt")])
        if f:
            self.ppt_path_var.set(f)
            self._log(f"파일: {os.path.basename(f)}")

    def _start_thread(self):
        if not self.ppt_path_var.get():
            return
        self.start_btn.config(state="disabled", bg=BG_BORDER, fg=FG_MUTED)
        threading.Thread(target=self._run, daemon=True).start()

    # ── 분석 로직 (python-pptx 크로스플랫폼) ─────────────────
    def _collect_texts_pptx(self, shape):
        results = []
        try:
            results.append(shape.name.lower())
            from pptx.util import Emu
            from pptx.enum.shapes import MSO_SHAPE_TYPE
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                for child in shape.shapes:
                    results.extend(self._collect_texts_pptx(child))
            else:
                if shape.has_text_frame:
                    t = shape.text_frame.text.strip().lower()
                    if t:
                        results.append(t)
        except Exception:
            pass
        return results

    def _scan_collection_pptx(self, shapes, keywords, slide_width_emu):
        for shape in shapes:
            try:
                texts = self._collect_texts_pptx(shape)
                for kw in keywords:
                    if any(kw in t for t in texts):
                        center = shape.left + shape.width / 2
                        pct = center / slide_width_emu * 100
                        if pct < 38:   return "left", pct
                        elif pct > 62: return "right", pct
                        else:          return "front", pct
            except Exception:
                continue
        return None

    def _get_instructor_pos_pptx(self, slide, keywords, slide_width_emu):
        try:
            r = self._scan_collection_pptx(slide.slide_layout.shapes, keywords, slide_width_emu)
            if r: return r
        except Exception:
            pass
        try:
            r = self._scan_collection_pptx(slide.slide_layout.slide_master.shapes, keywords, slide_width_emu)
            if r: return r
        except Exception:
            pass
        return None

    def _run(self):
        try:
            from pptx import Presentation
            p = self.ppt_path_var.get()
            self.log_text.delete(1.0, tk.END)
            self._log("분석 시작...")

            keywords = [k.strip().lower() for k in self.shape_keyword_var.get().split(",") if k.strip()]
            self._log(f"키워드: {keywords}")

            prs = Presentation(p)
            total_slides = len(prs.slides)
            slide_width_emu = prs.slide_width

            start_s = max(1, int(self.start_slide_var.get() or 1))
            end_s   = min(total_slides, int(self.end_slide_var.get() or 999))

            if start_s > end_s:
                raise Exception(f"시작({start_s}) > 끝({end_s})")

            self._log(f"범위: {start_s} ~ {end_s}  (총 {end_s - start_s + 1}장)\n")

            res_list  = []
            stats     = {"left": 0, "front": 0, "right": 0}
            kr_map    = {"front": "정면", "left": "좌측", "right": "우측"}
            not_found = []

            for i in range(start_s, end_s + 1):
                slide = prs.slides[i - 1]
                title = f"Slide {i}"
                try:
                    if slide.shapes.title:
                        title = slide.shapes.title.text.strip()
                except Exception:
                    pass

                result = self._get_instructor_pos_pptx(slide, keywords, slide_width_emu)
                if result:
                    angle, pct = result
                else:
                    angle, pct = "front", 50.0
                    not_found.append(i)

                stats[angle] += 1
                res_list.append({
                    "slide_number": i,
                    "title": title,
                    "angle": angle,
                    "position_pct": round(pct, 1),
                    "detected": result is not None
                })
                flag = "" if result else "  미감지"
                self._log(f"  #{i:03d}  {kr_map[angle]}  ({pct:.1f}%){flag}  {title[:35]}")

            out_json = p.replace(".pptx", "_분석.json").replace(".ppt", "_분석.json")
            with open(out_json, "w", encoding="utf-8") as f:
                json.dump({"ppt_file": p, "range": f"{start_s}-{end_s}", "slides": res_list},
                          f, ensure_ascii=False, indent=2)

            if not_found:
                self._log(f"\n미감지 슬라이드: {not_found}")
                self._log("  → 해당 슬라이드 마스터 도형 텍스트를 키워드에 추가하세요.")

            self._log(f"\n완료  →  {out_json}")
            self.root.after(0, lambda: self.stats_label.config(
                text=f"정면 {stats['front']}장  |  좌측 {stats['left']}장  |  우측 {stats['right']}장   "
                     f"(미감지 {len(not_found)}장)",
                fg=ACCENT))
            messagebox.showinfo("완료", f"분석 완료\n미감지: {len(not_found)}장")

        except Exception as e:
            self._log(f"오류: {e}")
            messagebox.showerror("오류", str(e))
        finally:
            self.start_btn.config(state="normal", bg=BTN_BG, fg=BTN_FG)


if __name__ == "__main__":
    import tempfile, atexit
    _LOCK_FILE = os.path.join(tempfile.gettempdir(), "SJE_PPT_Slide_Analyzer.lock")

    def _check_single_instance():
        if os.path.exists(_LOCK_FILE):
            try:
                with open(_LOCK_FILE, "r") as f:
                    pid = int(f.read().strip())
                # 해당 pid 프로세스가 살아있으면 중복 실행
                if platform.system() == "Windows":
                    import ctypes
                    handle = ctypes.windll.kernel32.OpenProcess(0x100000, False, pid)
                    alive = handle != 0
                    if handle:
                        ctypes.windll.kernel32.CloseHandle(handle)
                else:
                    import signal
                    os.kill(pid, 0)
                    alive = True
            except Exception:
                alive = False

            if alive:
                root_temp = tk.Tk()
                root_temp.withdraw()
                messagebox.showwarning("경고", "프로그램이 이미 실행 중입니다.")
                sys.exit(0)

        with open(_LOCK_FILE, "w") as f:
            f.write(str(os.getpid()))

        def _remove_lock():
            try:
                os.remove(_LOCK_FILE)
            except Exception:
                pass
        atexit.register(_remove_lock)

    _check_single_instance()

    root = tk.Tk()
    root.withdraw()  # 라이선스 확인 전까지 숨김

    lic = LicenseDialog(root)
    lic.check_and_run()

    if not lic.result:
        root.destroy()
        sys.exit(0)

    root.deiconify()
    app = PPTAnalyzerGUI(root, expiry_date=lic.expiry_date, plan=lic.plan, license_info=lic.license_info)
    root.mainloop()
